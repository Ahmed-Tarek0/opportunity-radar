from __future__ import annotations

import io
from pathlib import Path

from ..exceptions import CVParsingError
from ..logging_config import get_logger

logger = get_logger("cv_parser.text_extractor")

MIN_TEXT_LENGTH = 50  # characters minimum for a usable CV


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise CVParsingError(
            "PyMuPDF is not installed. Please install 'PyMuPDF' to parse PDF files."
        )

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages: list[str] = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                pages.append(text)
        doc.close()
        full_text = "\n\n".join(pages)
    except Exception as exc:
        raise CVParsingError(f"Failed to parse PDF: {exc}") from exc

    if len(full_text.strip()) < MIN_TEXT_LENGTH:
        raise CVParsingError(
            "Could not extract sufficient text from this PDF. "
            "The file may be scanned or image-based. "
            "Please upload a text-based PDF, or convert your CV to DOCX format."
        )

    logger.info("PDF text extracted", extra={"chars": len(full_text), "pages": len(pages)})
    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise CVParsingError(
            "python-docx is not installed. Please install 'python-docx' to parse DOCX files."
        )

    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)
        # Also extract table contents
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in paragraphs:
                        paragraphs.append(text)
        full_text = "\n\n".join(paragraphs)
    except Exception as exc:
        raise CVParsingError(f"Failed to parse DOCX: {exc}") from exc

    if len(full_text.strip()) < MIN_TEXT_LENGTH:
        raise CVParsingError(
            "Could not extract sufficient text from this DOCX file. "
            "The file may be empty or corrupted."
        )

    logger.info("DOCX text extracted", extra={"chars": len(full_text)})
    return full_text


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract text from PDF or DOCX based on file extension."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif suffix == ".docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise CVParsingError(f"Unsupported file type: {suffix}")
